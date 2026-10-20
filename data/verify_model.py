"""Verify the Excel model: formulas present + independent recomputation."""
from pathlib import Path

from openpyxl import load_workbook

import duckdb

wb = load_workbook(Path(__file__).parent.parent / "model" / "ValueAtStake.xlsx")
print("sheets:", wb.sheetnames)

m = wb["Model"]
for cell in ("B4", "B5", "B6", "B8", "B12", "B13", "B15", "B18", "B25"):
    print(f"Model!{cell}: {m[cell].value}")
s = wb["Scenarios"]
for cell in ("B4", "B5", "B6", "F4", "F5", "F6"):
    print(f"Scenarios!{cell}: {s[cell].value}")
sn = wb["Sensitivity"]
print("Sensitivity!B5:", sn["B5"].value)
print("Sensitivity!F9:", sn["F9"].value)

# ---- independent recomputation from the database ----
con = duckdb.connect(str(Path(__file__).parent / "org_effectiveness.duckdb"), read_only=True)
rows = con.execute(
    """
    SELECT company_name, role, fy,
           MAX(CASE WHEN metric='revenue_gbp' THEN value END) AS rev,
           MAX(CASE WHEN metric='staff_costs_gbp' THEN value END) AS staff,
           MAX(CASE WHEN metric='avg_headcount' THEN value END) AS hc
    FROM financials GROUP BY 1,2,3
    """
).fetchall()
latest = {}
for name, role, fy, rev, staff, hc in rows:
    if name not in latest or fy > latest[name][2]:
        latest[name] = [name, role, fy, rev, staff, hc]
allrows = list(latest.values())
peers = [r for r in allrows if r[1] == "peer"]
tgt = [r for r in allrows if r[1] == "target"][0]
import statistics

rev_fte = {r[0]: r[3] / r[5] for r in allrows}
staff_pct = {r[0]: r[4] / r[3] for r in allrows}
med_rf = statistics.median(r[3] / r[5] for r in peers)
med_sp = statistics.median(r[4] / r[3] for r in peers)
A = (tgt[4] / tgt[3] - med_sp) * tgt[3]
implied_hc = tgt[3] / med_rf
B = (tgt[5] - implied_hc) * (tgt[4] / tgt[5])
print(f"\nexpected: peer median rev/FTE {med_rf:.1f} | median staff% {med_sp * 100:.1f}%")
print(f"expected Method A {A:,.0f} k | Method B {B:,.0f} k | cross-check {abs(A - B) / A * 100:.1f}%")
print(f"target rev/FTE {rev_fte[tgt[0]]:.1f} vs median -> gap {(1 - rev_fte[tgt[0]] / med_rf) * 100:.1f}%")